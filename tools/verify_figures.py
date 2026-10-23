# -*- coding: utf-8 -*-
"""验证：说明md题号 ↔ xlsx J列 ↔ 图片文件 三者一致"""
import openpyxl, os, re

base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CH = ["第1章-命题逻辑","第2章-谓词逻辑","第3章-集合论","第4章-二元关系","第5章-函数","第6章-代数系统基础","第7章-图论基础"]

total_yes = 0
problems = []
for ch in CH:
    # 说明md 的题号 -> 图片文件
    md = os.path.join(base, "图形题配套包", ch, ch + "-图形题说明.md")
    txt = open(md, encoding="utf-8").read()
    md_map = {}   # 题号 -> 图片文件名
    for m in re.finditer(r"^\|\s*(题\d+)\s*\|\s*([^|]+?)\s*\|", txt, re.M):
        fig, img = m.group(1), m.group(2).strip()
        md_map[fig] = img
    # 从表格行解析 对应习题库题号（最后一列）
    md_qnos = {}   # 图片文件名 -> [题号]
    for m in re.finditer(r"^\|\s*(题\d+)\s*\|.*\|\s*([^|]*?)\s*\|\s*$", txt, re.M):
        fig = m.group(1); lastcol = m.group(2)
        img = md_map.get(fig)
        qnos = [int(x) for x in re.findall(r"第(\d+)题", lastcol)]
        md_qnos[img] = qnos
    md_yes = sorted({q for qs in md_qnos.values() for q in qs})

    # xlsx J列
    p = os.path.join(base, "习题库", ch + "-习题库.xlsx")
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb.worksheets[0]
    xlsx_yes = [r-1 for r in range(2, ws.max_row+1) if ws.cell(r, 10).value == "是"]
    wb.close()

    total_yes += len(xlsx_yes)
    # 一致性检查
    if md_yes != xlsx_yes:
        problems.append(f"{ch}: 说明md题号{md_yes} != xlsx J列{xlsx_yes}")
    # 图片文件存在性
    missing_img = [img for img in md_qnos if md_qnos[img] and not os.path.exists(os.path.join(base, "图形题配套包", ch, img))]
    if missing_img:
        problems.append(f"{ch}: 缺图片文件 {missing_img}")
    print(f"{ch}: J列=是 {len(xlsx_yes)} 题 | 说明md题号 {len(md_yes)} 题 | 图片映射 {len(md_qnos)} 张")

print(f"\n总图形题数: {total_yes}")
print("问题:", problems if problems else "无 —— 说明md / J列 / 图片文件 三者一致")
if problems:
    raise SystemExit(1)
