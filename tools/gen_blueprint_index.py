# -*- coding: utf-8 -*-
import os, re, docx, openpyxl, glob

base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
chapters = ["第1章-命题逻辑", "第2章-谓词逻辑", "第3章-集合论", "第4章-二元关系", "第5章-函数", "第6章-代数系统基础", "第7章-图论基础"]

# ---------- 1) 生成 考点分布.md（含「知识大点ID」列） ----------
# 87 条考点 → 知识大点ID 的全量映射，供智能体做「考点 → 大点 → 题目」精确路由。
# 该映射与 topic_map.md 大点表 / knowledge_dependency.md 一致，禁止改动 ID。
KPD_ID = {
    1: {
        "命题的概念与判断": "CH1-01",
        "命题符号化": "CH1-01",
        "联结词真值表与公式类型": "CH1-02",
        "基本等值式（德摩根律、蕴含等值式等）": "CH1-03",
        "等值演算化简与公式类型判断": "CH1-03",
        "析取范式与合取范式": "CH1-04",
        "主析取范式与主合取范式": "CH1-04",
        "极小项与极大项的编码": "CH1-04",
        "联结词完备集": "CH1-05",
        "推理的形式结构与推理定律": "CH1-06",
        "自然推理系统P构造证明": "CH1-06",
        "附加前提证明法与归谬法": "CH1-06",
    },
    2: {
        "个体词与谓词的概念": "CH2-01",
        "量词与命题符号化": "CH2-01",
        "自由变元与约束变元、辖域": "CH2-02",
        "换名规则与代替规则": "CH2-02",
        "解释与公式分类": "CH2-02",
        "量词否定等值式": "CH2-03",
        "量词辖域收缩与扩张": "CH2-03",
        "量词分配等值式": "CH2-03",
        "多个量词的顺序": "CH2-03",
        "前束范式": "CH2-04",
        "量词消去与引入规则（UI/UG/EI/EG）": "CH2-05",
        "自然推理系统构造证明": "CH2-05",
    },
    3: {
        "集合的概念与表示法": "CH3-01",
        "集合间关系与特殊集合（空集/幂集）": "CH3-01",
        "集合的基本运算（并交差补对称差）": "CH3-02",
        "集合运算律与恒等式证明": "CH3-02",
        "容斥原理（包含排斥原理）": "CH3-03",
        "笛卡尔积与有序对": "CH3-04",
        "二元关系的定义与表示": "CH3-04",
        "关系的运算（复合/逆/幂）": "CH3-04",
        "关系的性质（自反/对称/传递等）": "CH3-04",
        "关系的闭包（自反/对称/传递）": "CH3-04",
        "等价关系与划分、偏序关系与哈斯图": "CH3-05",
        "函数的定义与性质（单射/满射/双射）": "CH3-06",
    },
    4: {
        "有序对与笛卡尔积": "CH4-01",
        "二元关系的定义与表示（矩阵/图）": "CH4-02",
        "关系的运算（逆/复合/幂）": "CH4-03",
        "关系的性质（自反/对称/传递/反对称）": "CH4-04",
        "关系性质的运算保持": "CH4-04",
        "自反闭包与对称闭包": "CH4-05",
        "传递闭包": "CH4-05",
        "等价关系判定": "CH4-06",
        "等价类与商集": "CH4-06",
        "集合的划分与等价关系对应": "CH4-06",
        "偏序关系与哈斯图": "CH4-07",
        "偏序集特殊元素（最大/极小/上界/确界）": "CH4-07",
    },
    5: {
        "函数的定义与判定": "CH5-01",
        "函数的定义域与值域": "CH5-01",
        "单射的判定与证明": "CH5-02",
        "满射的判定与证明": "CH5-02",
        "双射的判定": "CH5-02",
        "复合函数运算与性质": "CH5-03",
        "反函数的存在条件与求解": "CH5-03",
        "函数计数（函数/单射/双射个数）": "CH5-04",
        "函数的限制与延拓": "CH5-03",
        "集合的基数与等势证明": "CH5-04",
        "可数集与不可数集": "CH5-04",
        "鸽巢原理及其应用": "CH5-05",
        "置换与递归函数": "CH5-05",
    },
    6: {
        "二元运算与封闭性": "CH6-01",
        "运算性质（交换/结合/分配律）": "CH6-01",
        "特异元素（单位元/零元/逆元）": "CH6-01",
        "半群与独异点判定": "CH6-02",
        "群的定义与判定": "CH6-03",
        "群的基本性质（消去律/逆元/穿脱）": "CH6-03",
        "子群判定": "CH6-03",
        "循环群与元素的阶": "CH6-03",
        "置换群与置换复合": "CH6-03",
        "环的定义与零因子": "CH6-04",
        "整环与域的判定": "CH6-04",
        "格的定义与运算（∧/∨）": "CH6-05",
        "布尔代数与德摩根律": "CH6-05",
    },
    7: {
        "图的基本概念与顶点度数": "CH7-01",
        "握手定理及推论": "CH7-01",
        "路与连通性（割点/桥）": "CH7-02",
        "图的矩阵表示（邻接矩阵A^k）": "CH7-03",
        "欧拉图判定": "CH7-04",
        "哈密顿图（必要条件/充分条件）": "CH7-04",
        "树的定义与性质": "CH7-05",
        "生成树与最小生成树": "CH7-05",
        "根树与二叉树性质": "CH7-05",
        "平面图与欧拉公式": "CH7-06",
        "平面图判定（K5/K3,3）": "CH7-06",
        "图的着色（点着色/边着色）": "CH7-07",
        "有向图连通性（强/单向/弱）": "CH7-02",
    },
}

lines = []
lines.append("---")
lines.append("doc_type: exam_blueprint")
lines.append("chapter: 全局")
lines.append("knowledge_ids: all")
lines.append("---")
lines.append("")
lines.append("# 离散数学考点分布表（考前冲刺蓝图）")
lines.append("")
lines.append("> 本表摘自各章《考点与易错点.docx》的「考点分布表」，用于智能体提供考前冲刺/高频考点规划与复习时间分配建议。")
lines.append("> 考察频率为「高」的考点优先复习；「中」需理解并能运用；「低」了解概念即可。")
lines.append("> 智能体可按考点频率+分值占比，为用户生成个性化冲刺计划，并结合 knowledge_dependency.md 检查薄弱前置。")
lines.append("> 每行末尾的「知识大点ID」用于「考点 → 大点 → 题目」精确路由（与 topic_map.md 大点表一致）。")
lines.append("")

for ch in chapters:
    p = os.path.join(base, "配套素材", ch + "-考点与易错点.docx")
    d = docx.Document(p)
    chnum = ch[1]  # "第1章..." -> "1"
    chno = int(chnum)
    lines.append(f"## 第{chnum}章（{ch[4:]}）")
    lines.append("")
    lines.append("| 核心考点 | 考察频率 | 常考题型 | 分值占比 | 知识大点ID |")
    lines.append("|---|---|---|---|---|")
    if d.tables:
        tbl = d.tables[0]
        for i, row in enumerate(tbl.rows):
            if i == 0:
                continue  # 表头
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if not any(cells):
                continue
            kp_name = cells[0].strip() if cells else ""
            kp_id = KPD_ID.get(chno, {}).get(kp_name, "")
            lines.append("| " + " | ".join(cells) + " | " + kp_id + " |")
    lines.append("")

open(os.path.join(base, "考点分布.md"), "w", encoding="utf-8").write("\n".join(lines))
print("考点分布.md done")

# ---------- 2) 重建 知识库元数据索引.xlsx（全量资产清单）----------
idx_path = os.path.join(base, "知识库元数据索引.xlsx")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "知识库索引"

# 资产类型判定
def classify(rel):
    ext = rel.rsplit(".", 1)[-1].lower() if "." in rel else "other"
    if rel.startswith("知识点/") and rel.endswith(".md"): return "知识点", "md"
    if rel.startswith("习题库/") and rel.endswith(".xlsx"): return "习题库", "xlsx"
    if rel.startswith("配套素材/") and rel.endswith(".docx"): return "考点与易错点", "docx"
    if rel.startswith("图形题配套包/") and rel.endswith(".md"): return "图形题说明", "md"
    if rel.startswith("图形题配套包/") and rel.endswith(".png"): return "图形题图片", "png"
    if rel.endswith("error_points_kb.md"): return "易错点库", "md"
    if rel.endswith("knowledge_dependency.md"): return "依赖表", "md"
    if rel.endswith("topic_map.md"): return "映射表", "md"
    if rel.endswith("考点分布.md"): return "考点蓝图", "md"
    if rel.endswith("高频考点清单.md"): return "高频考点", "md"
    if rel.endswith("题型归类映射表.md"): return "题型映射", "md"
    if rel == "AGENTS.md": return "智能体主配置", "md"
    if rel.endswith("student_progress.json"): return "进度状态", "json"
    if rel == "README.md": return "说明文档", "md"
    if rel.startswith("tools/"): return "维护脚本", ext
    if rel.startswith("参赛材料/参考资料/"): return "参考资料", ext
    if rel.startswith("参赛材料/"): return "参赛材料", ext
    if rel.startswith("skills/") and rel.endswith("SKILL.md"): return "技能定义", "md"
    return "其他", "other"

rows = []
for dirpath, dirnames, filenames in os.walk(base):
    dirnames[:] = [d for d in dirnames if d not in (".tmp", ".dumate", ".git", "__pycache__")]
    for fn in filenames:
        full = os.path.join(dirpath, fn)
        rel = os.path.relpath(full, base).replace("\\", "/")
        if rel.startswith("."):
            continue
        if rel == "知识库元数据索引.xlsx":  # 排除索引文件自身
            continue
        cat, ext = classify(rel)
        # 章节
        m = re.search(r"第(\d)章", rel)
        chapter = f"第{m.group(1)}章" if m else "全局"
        rows.append((rel, cat, chapter, ext))

rows.sort(key=lambda x: (x[1], x[0]))

# 表头
headers = ["序号", "相对路径", "资产类型", "所属章节", "文件格式", "状态"]
ws.append(headers)
for i, (rel, cat, chapter, ext) in enumerate(rows, 1):
    ws.append([i, rel, cat, chapter, ext, "完整"])

# 样式：表头加粗
from openpyxl.styles import Font
for c in ws[1]:
    c.font = Font(bold=True)
ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 52
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 10
ws.column_dimensions["E"].width = 10
ws.column_dimensions["F"].width = 8

wb.save(idx_path)
print("知识库元数据索引.xlsx done, 资产条数=", len(rows))
