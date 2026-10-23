# -*- coding: utf-8 -*-
"""核实图形题：说明md的映射 vs 习题库J列与题目内容"""
import openpyxl, os, re

base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CH = ["第1章-命题逻辑","第2章-谓词逻辑","第3章-集合论","第4章-二元关系","第5章-函数","第6章-代数系统基础","第7章-图论基础"]

out = []
for ch in CH:
    out.append("\n" + "="*70)
    out.append(ch)
    # 1) 说明md的映射
    md = os.path.join(base, "图形题配套包", ch, ch + "-图形题说明.md")
    mdtxt = open(md, encoding="utf-8").read()
    out.append("--- 说明md 映射（图片 -> 习题库题号）---")
    for m in re.finditer(r"^\|\s*(题\d+)\s*\|\s*([^|]+)\s*\|\s*([^|]+)\s*\|\s*([^|]+)\s*\|\s*([^|]+)\s*\|", mdtxt, re.M):
        fig, img, desc, kp, qnos = m.groups()
        out.append(f"  {fig} [{img.strip()}] 知识点[{kp.strip()}] -> {qnos.strip()}")
    # 2) xlsx 中 J列=是 的题目
    p = os.path.join(base, "习题库", ch + "-习题库.xlsx")
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb.worksheets[0]
    out.append("--- 习题库 J列=是 的题目 ---")
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 10).value == "是":
            qno = r - 1
            qid = ws.cell(r, 8).value
            qtext = str(ws.cell(r, 1).value)[:60]
            kp = ws.cell(r, 5).value
            out.append(f"  第{qno}题 [{qid}] {kp} | {qtext}")
    wb.close()

open(os.path.join(base, "tools", "fig_audit.txt"), "w", encoding="utf-8").write("\n".join(out))
print("done, lines:", len(out))
