# -*- coding: utf-8 -*-
"""修复图形题映射：按知识点精确重映射 图片->题目，修正 J列 与 说明md 的题号列"""
import openpyxl, os, re

base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# 章节 -> {图片题号: [正确对应的习题库题号]}
# 依据：图片"对应知识点" 与 题目"对应知识点(E列)" 精确匹配
FIGMAP = {
    "第1章-命题逻辑":    {1: [3], 2: [], 3: [7], 4: [], 5: []},
    "第2章-谓词逻辑":    {1: [1], 2: [2], 3: [3], 4: [5], 5: [7]},
    "第3章-集合论":      {1: [2], 2: [2], 3: [5], 4: [5], 5: [14], 6: [12], 7: [13], 8: [14, 25]},
    "第4章-二元关系":    {1: [2], 2: [4], 3: [5], 4: [3], 5: [13], 6: [15], 7: [16], 8: [17], 9: [21], 10: [25]},
    "第5章-函数":        {1: [1], 2: [3], 3: [4], 4: [11], 5: [5], 6: [6]},
    "第6章-代数系统基础": {1: [11, 12], 2: [5], 3: [3], 4: [3], 5: [4], 6: [14]},
    "第7章-图论基础":    {1: [1, 11], 2: [2], 3: [3], 4: [4, 11], 5: [5, 15], 6: [6, 13], 7: [7],
                          8: [8, 12, 24], 9: [20], 10: [14], 11: [9, 19], 12: [10]},
}

def qnos_to_str(qs):
    if not qs:
        return "无（备用插图）"
    return "、".join(f"第{q}题" for q in qs)

report = []
for ch, mapping in FIGMAP.items():
    # 该章所有应为"是"的题号集合
    yes = sorted({q for qs in mapping.values() for q in qs})

    # ---- 1) 修正 xlsx J列 ----
    p = os.path.join(base, "习题库", ch + "-习题库.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    for r in range(2, ws.max_row + 1):
        qno = r - 1
        ws.cell(r, 10, value=("是" if qno in yes else "否"))
    wb.save(p)
    wb.close()

    # ---- 2) 修正 说明md 的"对应习题库题号"列 ----
    md = os.path.join(base, "图形题配套包", ch, ch + "-图形题说明.md")
    txt = open(md, encoding="utf-8").read()
    lines = txt.splitlines()
    out = []
    for line in lines:
        m = re.match(r"^\|\s*(题\d+)\s*\|", line)
        if m:
            fig = m.group(1)
            num = int(fig.replace("题", ""))
            qs = mapping.get(num, [])
            newval = qnos_to_str(qs)
            # 替换行末最后一列（"| ... |"结尾）
            line = re.sub(r"\|\s*[^|]*\|\s*$", f"| {newval} |", line)
        out.append(line)
    newtxt = "\n".join(out)
    # 修复 ch3 描述中的 |A| 等竖线（破坏markdown表格）
    newtxt = newtxt.replace("|A|", "｜A｜").replace("|B|", "｜B｜")
    newtxt = newtxt.replace("|A∩B|", "｜A∩B｜").replace("|A∪B|", "｜A∪B｜").replace("|A∪B∪C|", "｜A∪B∪C｜")
    open(md, "w", encoding="utf-8").write(newtxt)

    report.append(f"{ch}: J列=是 {len(yes)} 题 {yes}")

print("\n".join(report))
print("\n总图形题数:", sum(len(set(q for qs in m.values() for q in qs)) for m in FIGMAP.values()))
